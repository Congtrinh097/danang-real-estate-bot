import io
import pytest
from docx import Document
from pypdf import PdfWriter
from backend.document_converter import DocumentConverter

def test_convert_markdown():
    sample_md = "# Luật Nhà Ở 2023\n## Chương I\n### Điều 1. Phạm vi điều chỉnh\nNội dung điều 1."
    md_bytes = sample_md.encode("utf-8")
    result = DocumentConverter.convert_to_markdown(md_bytes, "luat_nha_o.md")
    assert "# Luật Nhà Ở 2023" in result
    assert "### Điều 1. Phạm vi điều chỉnh" in result

def test_convert_docx():
    doc = Document()
    doc.add_heading("Luật Nhà Ở 2023", level=1)
    doc.add_heading("Chương II: Sở hữu nhà ở", level=2)
    doc.add_heading("Điều 5: Quyền sở hữu", level=3)
    doc.add_paragraph("Khoản 1: Công dân Việt Nam có quyền sở hữu nhà ở hợp pháp.")
    
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docxio = docx_io.getvalue()

    result = DocumentConverter.convert_to_markdown(docx_bytes, "sample.docx")
    assert "# Luật Nhà Ở 2023" in result
    assert "## Chương II: Sở hữu nhà ở" in result
    assert "### Điều 5: Quyền sở hữu" in result
    assert "Khoản 1: Công dân Việt Nam" in result

def test_convert_empty_file():
    with pytest.raises(ValueError, match="rỗng"):
        DocumentConverter.convert_to_markdown(b"", "empty.pdf")

def test_convert_pdf(monkeypatch):
    from pypdf import PageObject
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    
    pdf_io = io.BytesIO()
    writer.write(pdf_io)
    pdf_bytes = pdf_io.getvalue()
    
    monkeypatch.setattr(PageObject, "extract_text", lambda self: "Chương I: Quy Định Chung\nĐiều 1: Phạm vi áp dụng luật")
    result = DocumentConverter.convert_to_markdown(pdf_bytes, "sample.pdf")
    assert "## Chương I: Quy Định Chung" in result
    assert "### Điều 1: Phạm vi áp dụng luật" in result

def test_convert_blank_pdf(monkeypatch):
    from pypdf import PageObject
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    
    pdf_io = io.BytesIO()
    writer.write(pdf_io)
    pdf_bytes = pdf_io.getvalue()

    monkeypatch.setattr(PageObject, "extract_text", lambda self: "")
    with pytest.raises(ValueError, match="Không trích xuất được nội dung"):
        DocumentConverter.convert_to_markdown(pdf_bytes, "blank.pdf")

def test_convert_unsupported_file_format():
    with pytest.raises(ValueError, match="Định dạng file không được hỗ trợ"):
        DocumentConverter.convert_to_markdown(b"binarycontent", "malware.exe")
