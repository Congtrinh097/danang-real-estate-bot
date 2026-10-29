import io
import os
from typing import Optional
from pypdf import PdfReader
from docx import Document

class DocumentConverter:
    """Chuyển đổi các định dạng tài liệu (PDF, Word DOCX, Markdown, Text) sang định dạng Markdown chuẩn hóa."""

    @staticmethod
    def convert_to_markdown(file_bytes: bytes, filename: str) -> str:
        if not file_bytes or len(file_bytes) == 0:
            raise ValueError("File tải lên rỗng (0 bytes). Vui lòng chọn file có nội dung.")

        ext = os.path.splitext(filename)[1].lower()
        
        if ext in [".md", ".markdown", ".txt"]:
            return DocumentConverter._convert_text_or_markdown(file_bytes)
        elif ext in [".docx"]:
            return DocumentConverter._convert_docx(file_bytes)
        elif ext in [".pdf"]:
            return DocumentConverter._convert_pdf(file_bytes)
        else:
            raise ValueError(f"Định dạng file không được hỗ trợ ({ext}). Chỉ hỗ trợ .pdf, .docx, .md, .txt")

    @staticmethod
    def _convert_text_or_markdown(file_bytes: bytes) -> str:
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1", errors="ignore")

    @staticmethod
    def _convert_docx(file_bytes: bytes) -> str:
        doc = Document(io.BytesIO(file_bytes))
        md_lines = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower() if p.style else ""
            if "heading 1" in style_name:
                md_lines.append(f"# {text}\n")
            elif "heading 2" in style_name:
                md_lines.append(f"## {text}\n")
            elif "heading 3" in style_name:
                md_lines.append(f"### {text}\n")
            else:
                md_lines.append(f"{text}\n")

        for table in doc.tables:
            table_md = []
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                table_md.append("| " + " | ".join(row_cells) + " |")
            if table_md:
                # Thêm header separator sau hàng đầu tiên
                header_sep = "| " + " | ".join(["---"] * len(table.rows[0].cells)) + " |"
                table_md.insert(1, header_sep)
                md_lines.append("\n" + "\n".join(table_md) + "\n")

        return "\n".join(md_lines)

    @staticmethod
    def _convert_pdf(file_bytes: bytes) -> str:
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
        except Exception as e:
            raise ValueError(f"Không thể đọc file PDF (file lỗi hoặc định dạng không hợp lệ): {str(e)}")

        pages_text = []
        for idx, page in enumerate(reader.pages):
            try:
                text = page.extract_text()
                if text:
                    pages_text.append(text.strip())
            except Exception:
                continue

        raw_text = "\n\n".join(pages_text)
        if not raw_text.strip():
            raise ValueError("Không trích xuất được nội dung văn bản từ file PDF (file scan dạng ảnh hoặc trang trắng). Vui lòng cung cấp file PDF chứa text hoặc định dạng Word/Markdown.")
        
        # Chuẩn hóa các dòng tiêu đề Điều / Chương nếu có
        lines = raw_text.split("\n")
        formatted_lines = []
        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue
            lower_str = line_str.lower()
            if lower_str.startswith("chương ") or lower_str.startswith("mục "):
                formatted_lines.append(f"\n## {line_str}\n")
            elif lower_str.startswith("điều ") or lower_str.startswith("dieu "):
                formatted_lines.append(f"\n### {line_str}\n")
            else:
                formatted_lines.append(line_str)

        return "\n".join(formatted_lines)
